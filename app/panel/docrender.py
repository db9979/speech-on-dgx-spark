"""Dokumente erstellen (plan plaene/dokumente-erstellen.md): the files themselves, built by fixed code.

The language model never writes a file. It fills a checked content plan (docmake.check_content: title,
sections with paragraphs, bullet points, one table, one chart, pictures of the profile's own documents,
speaker notes; for a letter the letter's parts) and this module turns that plan into Word (.docx),
PowerPoint (.pptx), Excel (.xlsx), PDF, Markdown, plain text or CSV with fixed layouts (STYLES) or a
profile's own template (docx/pptx, checked by docmake when it was uploaded).

What can never end up in a file: macros (no .docm/.pptm/.xlsm is ever written, templates with macros are
refused on upload), links to other files or the internet (no hyperlinks are written at all), formulas
other than SUMME/MITTELWERT/MIN/MAX/ANZAHL over a cell range (everything else is written as text, so a
cell "=HYPERLINK(...)" stays a word), and in CSV a cell that starts with = + - @ gets a ' in front.
Pictures come only from paths docmake resolved inside the profile's own folder.

Pure functions over the content plan: render(content, fmt, style, pictures) -> bytes. No network, no
model, no clock (the date is part of the content). Runs in a worker thread (docmake), a few seconds at most.
"""
import csv
import io
import math
import os
import re

FORMATS = ("docx", "pptx", "xlsx", "pdf", "md", "txt", "csv")
CTYPE = {"docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
         "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
         "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
         "pdf": "application/pdf", "md": "text/markdown; charset=utf-8", "txt": "text/plain; charset=utf-8",
         "csv": "text/csv; charset=utf-8"}

# the built-in looks (Vorlagen): colors as hex without #; font for the text, head for titles and headings. Aptos is
# Office's own font (Windows, Mac, web); other programs put in a similar one. Georgia and Arial are everywhere.
STYLES = {
    "spark-hell": {"de": "Spark hell", "en": "Spark light", "accent": "5B5BF0", "accent2": "0EA5C6", "bg": "FFFFFF",
                   "fg": "141824", "muted": "5F6B85", "font": "Aptos", "head": "Aptos Display", "dark": False},
    "spark-dunkel": {"de": "Spark dunkel", "en": "Spark dark", "accent": "8B8CFF", "accent2": "22D3EE", "bg": "0B0D14",
                     "fg": "E8EBF2", "muted": "8B95AB", "font": "Aptos", "head": "Aptos Display", "dark": True},
    "frisch": {"de": "Frisch", "en": "Fresh", "accent": "0F766E", "accent2": "F59E0B", "bg": "FFFFFF", "fg": "10201E",
               "muted": "52645F", "font": "Aptos", "head": "Aptos Display", "dark": False},
    "warm": {"de": "Warm", "en": "Warm", "accent": "C2410C", "accent2": "7C3AED", "bg": "FFFBF5", "fg": "2A1A12",
             "muted": "7A5F52", "font": "Aptos", "head": "Georgia", "dark": False},
    "schlicht": {"de": "Schlicht", "en": "Plain", "accent": "2F3A4F", "accent2": "6B7A99", "bg": "FFFFFF", "fg": "1D2433",
                 "muted": "5F6B85", "font": "Calibri", "head": "Georgia", "dark": False},
    "brief": {"de": "Brief (DIN 5008)", "en": "Letter (DIN 5008)", "accent": "1D2433", "accent2": "5F6B85", "bg": "FFFFFF",
              "fg": "1D2433", "muted": "5F6B85", "font": "Arial", "head": "Arial", "dark": False},
}
PALETTE = ("EA8034", "16A34A", "DC2626", "9333EA", "0284C7", "CA8A04")   # after accent and accent2 (charts)
# fonts for PDF and chart pictures: the first that exists (DejaVu is on Ubuntu and DGX OS); without one the PDF
# uses the built-in Helvetica and characters outside Latin-1 are replaced (LATIN1)
FONT_DIRS = ("/usr/share/fonts/truetype/dejavu", "/usr/share/fonts/dejavu", "/usr/share/fonts/TTF")
LATIN1 = {"„": '"', "“": '"', "”": '"', "‚": "'", "‘": "'", "’": "'", "–": "-", "—": "-", "…": "...", "€": "EUR",
          "•": "-", "→": "->", "←": "<-", "≈": "~", "≤": "<=", "≥": ">=", "×": "x", "\u00a0": " ", "\u202f": " "}
FORMULA = re.compile(r"=\s*(SUMME|SUM|MITTELWERT|AVERAGE|MIN|MAX|ANZAHL|COUNT)\(\s*([A-Z]{1,2}[1-9][0-9]{0,4})\s*:"
                     r"\s*([A-Z]{1,2}[1-9][0-9]{0,4})\s*\)\s*", re.I)
EN_FN = {"SUMME": "SUM", "MITTELWERT": "AVERAGE", "ANZAHL": "COUNT"}
NUMBER_DE = re.compile(r"-?\d{1,3}(?:\.\d{3})+(?:,\d+)?|-?\d+(?:,\d+)?")
NUMBER_EN = re.compile(r"-?\d{1,3}(?:,\d{3})+(?:\.\d+)?|-?\d+(?:\.\d+)?")
MONTHS = ("Januar", "Februar", "März", "April", "Mai", "Juni", "Juli", "August", "September", "Oktober", "November",
          "Dezember")


def style_of(key):
    st = dict(STYLES.get(key) or STYLES["spark-hell"], key=key if key in STYLES else "spark-hell")
    # the light fill of boxes and cards, the line around them, the zebra rows of tables
    st["tint"] = mix(st["accent"], st["bg"], 0.80 if st["dark"] else 0.90)
    st["card"] = mix(st["fg"], st["bg"], 0.90 if st["dark"] else 0.965)
    st["line"] = mix(st["fg"], st["bg"], 0.78 if st["dark"] else 0.86)
    st["zebra"] = mix(st["accent"], st["bg"], 0.88 if st["dark"] else 0.95)
    return st


def _rgb(h):
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))


def mix(a, b, t):
    """Color a moved the share t towards color b (hex without #)."""
    x, y = _rgb(a), _rgb(b)
    return "".join(f"{round(p + (q - p) * t):02X}" for p, q in zip(x, y))


def chart_colors(style):
    return [style["accent"], style["accent2"]] + list(PALETTE)


def number(s, lang="de"):
    """A cell that is only a number ("1.234,5", "12", "-3,2"; English "1,234.5") -> float, else None."""
    s = str(s).strip()
    rx = NUMBER_EN if lang == "en" else NUMBER_DE
    if not s or len(s) > 24 or not rx.fullmatch(s):
        return None
    if lang == "en":
        s = s.replace(",", "")
    else:
        s = s.replace(".", "").replace(",", ".")
    try:
        v = float(s)
    except ValueError:
        return None
    return v if math.isfinite(v) else None


def formula(s):
    """The English formula for an allowed one (=SUMME(B2:B9) -> =SUM(B2:B9)), else None."""
    m = FORMULA.fullmatch(str(s))
    if not m:
        return None
    fn = m.group(1).upper()
    return f"={EN_FN.get(fn, fn)}({m.group(2).upper()}:{m.group(3).upper()})"


def csv_cell(s):
    """A cell for CSV: one that a spreadsheet would read as a formula gets a ' in front (CSV injection)."""
    s = str(s)
    return "'" + s if s[:1] in ("=", "+", "-", "@", "\t", "\r") and number(s) is None else s


def _font_files():
    for d in FONT_DIRS:
        reg, bold = os.path.join(d, "DejaVuSans.ttf"), os.path.join(d, "DejaVuSans-Bold.ttf")
        if os.path.isfile(reg) and os.path.isfile(bold):
            return reg, bold
    return None, None


def _sections(c):
    return c.get("sections") or []


def _all_pictures(c, pictures):
    return [p for s in _sections(c) for p in (pictures(s) if pictures else [])]


# ---------------------------------------------------------------- chart as a picture (Word, PDF)
def _nice(hi):
    """A round step for about 4 to 6 grid lines up to hi."""
    if hi <= 0:
        return 1.0
    raw = hi / 5
    mag = 10 ** math.floor(math.log10(raw))
    for m in (1, 2, 2.5, 5, 10):
        if raw <= m * mag:
            return m * mag
    return 10 * mag


def chart_png(chart, style, width=1200, height=640):
    """A bar, line or ring chart drawn with Pillow (Word and PDF have no own charts here): round grid steps, the
    values on the bars, the style's colors; drawn twice the size and scaled down for clean edges."""
    from PIL import Image, ImageDraw, ImageFont
    k = 2
    W_, H_ = width * k, height * k
    bg = _rgb(style["bg"]) if style["dark"] else (255, 255, 255)
    fg, mut, grid = _rgb(style["fg"]), _rgb(style["muted"]), _rgb(style.get("line") or "D0D5DE")
    cols = [_rgb(h) for h in chart_colors(style)]
    im = Image.new("RGB", (W_, H_), bg)
    d = ImageDraw.Draw(im)
    reg, bold = _font_files()
    try:
        f = ImageFont.truetype(reg, 22 * k) if reg else ImageFont.load_default(22 * k)
        fs = ImageFont.truetype(reg, 19 * k) if reg else ImageFont.load_default(19 * k)
        fb = ImageFont.truetype(bold, 30 * k) if bold else ImageFont.load_default(30 * k)
        fv = ImageFont.truetype(bold, 20 * k) if bold else ImageFont.load_default(20 * k)
    except (OSError, TypeError):
        f = fs = fb = fv = ImageFont.load_default()
    top = 24 * k
    if chart.get("title"):
        d.text((30 * k, top), chart["title"], fill=fg, font=fb)
        top += 56 * k
    labels, series = chart.get("labels") or [], chart.get("series") or []

    def done():
        return _png(im.resize((width, height), Image.LANCZOS))
    if not labels or not series:
        return done()
    if len(series) > 1 or chart.get("type") == "pie":
        lx = 30 * k
        names = [s.get("label") or "" for s in series] if chart.get("type") != "pie" else []
        for i, nm in enumerate(names):     # legend
            d.rounded_rectangle((lx, top + 6 * k, lx + 18 * k, top + 24 * k), 4 * k, fill=cols[i % len(cols)])
            d.text((lx + 26 * k, top), nm, fill=mut, font=f)
            lx += 48 * k + int(d.textlength(nm, font=f))
        if names:
            top += 46 * k
    if chart.get("type") == "pie":
        vals = [max(0.0, float(v)) for v in (series[0].get("values") or [])][:len(labels)]
        total = sum(vals) or 1.0
        size = min(H_ - top - 30 * k, 520 * k)
        box = (60 * k, top, 60 * k + size, top + size)
        a = -90.0
        for i, v in enumerate(vals):
            sweep = 360.0 * v / total
            d.pieslice(box, a, a + sweep, fill=cols[i % len(cols)], outline=bg, width=4 * k)
            a += sweep
        inner = size * 0.28
        cx, cy = (box[0] + box[2]) / 2, (box[1] + box[3]) / 2
        d.ellipse((cx - inner, cy - inner, cx + inner, cy + inner), fill=bg)     # a ring, not a pie
        y = top + 10 * k
        for i, lab in enumerate(labels[:len(vals)]):
            d.rounded_rectangle((box[2] + 50 * k, y + 6 * k, box[2] + 70 * k, y + 26 * k), 5 * k, fill=cols[i % len(cols)])
            d.text((box[2] + 82 * k, y), f"{lab}", fill=fg, font=f)
            pct = f"{100 * vals[i] / total:.0f} %"
            d.text((W_ - 40 * k - d.textlength(pct, font=fv), y), pct, fill=mut, font=fv)
            y += 42 * k
        return done()
    allv = [float(v) for s in series for v in (s.get("values") or [])] or [0.0]
    lo, hi = min(0.0, min(allv)), max(allv)
    step_v = _nice(max(hi - lo, 1e-9))
    lo, hi = math.floor(lo / step_v) * step_v, math.ceil(hi / step_v) * step_v
    if hi == lo:
        hi = lo + step_v
    x0, x1, y0, y1 = 110 * k, W_ - 30 * k, top + 30 * k, H_ - 60 * k
    v = lo
    while v <= hi + step_v / 2:     # grid with round values
        y = y1 - (y1 - y0) * (v - lo) / (hi - lo)
        d.line((x0, y, x1, y), fill=grid, width=k)
        lab = _fmt(v)
        d.text((x0 - 16 * k - d.textlength(lab, font=fs), y - 12 * k), lab, fill=mut, font=fs)
        v += step_v
    n = len(labels)
    step = (x1 - x0) / max(1, n)

    def ypos(v_):
        return y1 - (y1 - y0) * (float(v_) - lo) / (hi - lo)
    for i, lab in enumerate(labels):
        txt_ = str(lab)[:14]
        d.text((x0 + step * i + step / 2 - d.textlength(txt_, font=fs) / 2, y1 + 14 * k), txt_, fill=mut, font=fs)
    one = len(series) == 1 and n <= 14
    for si, s in enumerate(series):
        vals = (s.get("values") or [])[:n]
        col_ = cols[si % len(cols)]
        if chart.get("type") == "line":
            pts = [(x0 + step * i + step / 2, ypos(v_)) for i, v_ in enumerate(vals)]
            if len(pts) > 1:
                d.line(pts, fill=col_, width=6 * k, joint="curve")
            for p in pts:
                d.ellipse((p[0] - 9 * k, p[1] - 9 * k, p[0] + 9 * k, p[1] + 9 * k), fill=bg, outline=col_, width=4 * k)
        else:
            bw = step * 0.62 / max(1, len(series))
            for i, v_ in enumerate(vals):
                bx = x0 + step * i + step * 0.19 + bw * si
                ya, yb = min(ypos(v_), ypos(0)), max(ypos(v_), ypos(0))
                if yb - ya >= 2:
                    d.rounded_rectangle((bx, ya, bx + bw - 4 * k, yb), min(10 * k, (yb - ya) / 2, bw / 4), fill=col_)
                if one:
                    lab = _fmt(v_)
                    d.text((bx + (bw - 4 * k) / 2 - d.textlength(lab, font=fv) / 2, ya - 30 * k), lab, fill=fg, font=fv)
    return done()


def _fmt(v):
    v = float(v)
    return str(int(v)) if v == int(v) and abs(v) < 1e12 else f"{v:.2f}".rstrip("0").rstrip(".")


def _png(im):
    b = io.BytesIO()
    im.save(b, "PNG", optimize=True)
    return b.getvalue()


# ---------------------------------------------------------------- Word
def docx(c, style, pictures=None, base=None):
    if not base and not c.get("letter"):
        return _docx_designed(c, style, pictures)
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor
    doc = Document(base) if base else Document()
    if base:     # an own template: its styles, headers and footers stay, its body text goes
        body = doc.element.body
        for el in list(body):
            if not el.tag.endswith("}sectPr"):
                body.remove(el)
    else:
        st = doc.styles["Normal"]
        st.font.name = style["font"]
        st.font.size = Pt(11)
        st.font.color.rgb = RGBColor.from_string(style["fg"] if not style["dark"] else "141824")
        for name, size in (("Title", 26), ("Heading 1", 16), ("Heading 2", 13)):
            try:
                s = doc.styles[name]
                s.font.name = style["font"]
                s.font.size = Pt(size)
                s.font.color.rgb = RGBColor.from_string(style["accent"] if not style["dark"] else "2F3A4F")
            except KeyError:
                pass
        sec = doc.sections[0]
        sec.left_margin, sec.right_margin = Cm(2.5), Cm(2.0)
    if c.get("letter"):
        _docx_letter(doc, c, Pt, Cm, WD_ALIGN_PARAGRAPH)
    else:
        doc.add_paragraph(c.get("title") or "", style="Title")
        if c.get("subtitle"):
            p = doc.add_paragraph(c["subtitle"])
            p.runs[0].italic = True
        if c.get("date") or c.get("author"):
            p = doc.add_paragraph(" · ".join(x for x in (c.get("author"), c.get("date")) if x))
            p.runs[0].font.size = Pt(9)
        for s in _sections(c):
            if s.get("heading"):
                doc.add_heading(s["heading"], level=1)
            for t in s.get("paras") or []:
                doc.add_paragraph(t)
            for t in s.get("bullets") or []:
                doc.add_paragraph(t, style="List Bullet")
            if s.get("table"):
                _docx_table(doc, s["table"])
            if s.get("chart"):
                doc.add_picture(io.BytesIO(chart_png(s["chart"], style)), width=Cm(15))
            for p in (pictures(s) if pictures else []):
                doc.add_picture(p, width=Cm(12))
            if s.get("notes"):
                p = doc.add_paragraph(s["notes"])
                p.runs[0].italic = True
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _print_style(style):
    """Paper is white: a dark look prints with the light look's colors."""
    return style_of("spark-hell") if style["dark"] else style


KIND_LABEL = {"bericht": ("BERICHT", "REPORT"), "protokoll": ("PROTOKOLL", "MINUTES"), "notiz": ("NOTIZ", "NOTE"),
              "zusammenfassung": ("ZUSAMMENFASSUNG", "SUMMARY"), "einladung": ("EINLADUNG", "INVITATION"),
              "praesentation": ("PRÄSENTATION", "PRESENTATION"), "tabelle": ("TABELLE", "SPREADSHEET")}


def _docx_designed(c, style, pictures=None):
    """The Spark's own Word layout: a colored title block, numbered headings with a line, key numbers as cards, a
    core sentence in a box, tables with a colored head and zebra rows, page numbers in the footer."""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    st = _print_style(style)
    rgb = {k: RGBColor.from_string(st[k]) for k in ("accent", "accent2", "fg", "muted")}
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)     # A4
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(2.3), Cm(2.0), Cm(2.0), Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size, normal.font.color.rgb = st["font"], Pt(11), rgb["fg"]
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), st["font"])
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size in (("Heading 1", 17), ("Heading 2", 13)):
        try:
            h = doc.styles[name]
            h.font.name, h.font.size, h.font.bold, h.font.color.rgb = st["head"], Pt(size), True, rgb["fg"]
            rf = h.element.rPr.rFonts
            for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(a), st["head"])
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rf.get(qn(a)) is not None:
                    del rf.attrib[qn(a)]
            h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(18), Pt(6)
        except KeyError:
            pass
    width = sec.page_width - sec.left_margin - sec.right_margin

    def shade(cell, hex_):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_)
        tcPr.append(shd)

    def borders(cell, **edges):
        """edges: left=("single", 24, "5B5BF0"), ... in eighths of a point; unnamed edges get none."""
        tcPr = cell._tc.get_or_add_tcPr()
        b = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right"):
            e = OxmlElement(f"w:{edge}")
            kind, sz, color = edges.get(edge, ("nil", 0, "auto"))
            e.set(qn("w:val"), kind)
            if kind != "nil":
                e.set(qn("w:sz"), str(sz))
                e.set(qn("w:color"), color)
                e.set(qn("w:space"), "0")
            b.append(e)
        tcPr.append(b)

    def margins(cell, top, side):
        tcPr = cell._tc.get_or_add_tcPr()
        m = OxmlElement("w:tcMar")
        for edge, v in (("top", top), ("bottom", top), ("left", side), ("right", side)):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:w"), str(int(v)))
            e.set(qn("w:type"), "dxa")
            m.append(e)
        tcPr.append(m)

    def no_borders(table):
        tblPr = table._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "nil")
            b.append(e)
        tblPr.append(b)

    def run(p, value, size=None, bold=False, color=None, italic=False, font=None):
        r = p.add_run(value)
        if size:
            r.font.size = Pt(size)
        r.bold, r.italic = bold, italic
        if color is not None:
            r.font.color.rgb = color
        if font:
            r.font.name = font
            r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
        return r

    def box_table(cols):
        t = doc.add_table(rows=1, cols=cols)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        no_borders(t)
        trPr = t.rows[0]._tr.get_or_add_trPr()      # a box never breaks across two pages
        trPr.append(OxmlElement("w:cantSplit"))
        t.autofit = False
        for cell in t.rows[0].cells:
            cell.width = int(width / cols)
        return t

    en = c.get("lang") == "en"
    white = RGBColor(0xFF, 0xFF, 0xFF)
    # ---- the title block
    t = box_table(1)
    cell = t.rows[0].cells[0]
    shade(cell, st["accent"])
    margins(cell, 380, 420)
    p = cell.paragraphs[0]
    label = KIND_LABEL.get(c.get("kind"), ("", ""))[1 if en else 0]
    if label:
        run(p, label, 9, True, RGBColor.from_string(mix(st["accent"], "FFFFFF", 0.7)))
        p.paragraph_format.space_after = Pt(4)
        p = cell.add_paragraph()
    run(p, c.get("title") or "", 28 if len(c.get("title") or "") <= 50 else 22, True, white, font=st["head"])
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.0
    if c.get("subtitle"):
        p = cell.add_paragraph()
        run(p, c["subtitle"], 14, False, white)
        p.paragraph_format.space_after = Pt(2)
    info = " · ".join(x for x in (c.get("author"), c.get("date")) if x)
    if info:
        p = cell.add_paragraph()
        run(p, info, 9.5, False, RGBColor.from_string(mix(st["accent"], "FFFFFF", 0.75)))
        p.paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ---- the parts
    for k, s in enumerate(_sections(c), 1):
        if s.get("heading"):
            h = doc.add_heading("", level=1)
            run(h, f"{k:02d}  ", None, True, rgb["accent2"])
            run(h, s["heading"], None, True, rgb["fg"])
            pPr = h._p.get_or_add_pPr()
            pb = OxmlElement("w:pBdr")
            bot = OxmlElement("w:bottom")
            for a, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "4"), ("w:color", st["accent"])):
                bot.set(qn(a), v)
            pb.append(bot)
            pPr.append(pb)
        for x in s.get("paras") or []:
            doc.add_paragraph(x)
        if s.get("stats"):
            t = box_table(len(s["stats"]))
            for i, x in enumerate(s["stats"]):
                cell = t.rows[0].cells[i]
                shade(cell, st["tint"])
                borders(cell, left=("single", 24, st["accent"] if i % 2 == 0 else st["accent2"]))
                margins(cell, 160, 220)
                p = cell.paragraphs[0]
                run(p, x["value"], 22, True, rgb["accent"], font=st["head"])
                p.paragraph_format.space_after = Pt(0)
                if x.get("label"):
                    p = cell.add_paragraph()
                    run(p, x["label"], 9.5, False, rgb["fg"])
                    p.paragraph_format.space_after = Pt(0)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        for x in s.get("bullets") or []:
            doc.add_paragraph(x, style="List Bullet")
        if s.get("quote"):
            t = box_table(1)
            cell = t.rows[0].cells[0]
            shade(cell, st["tint"])
            borders(cell, left=("single", 36, st["accent"]))
            margins(cell, 200, 300)
            p = cell.paragraphs[0]
            run(p, "„" if not en else "“", 20, True, rgb["accent"], font=st["head"])
            run(p, s["quote"] + ("“" if not en else "”"), 13, True, rgb["fg"], font=st["head"])
            p.paragraph_format.space_after = Pt(0)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        if s.get("table"):
            _docx_table_designed(doc, s["table"], st, shade, borders, margins, no_borders, width)
        if s.get("chart"):
            doc.add_picture(io.BytesIO(chart_png(s["chart"], st)), width=min(width, Cm(16)))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for pic in (pictures(s) if pictures else []):
            doc.add_picture(pic, width=Cm(12))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if s.get("notes"):
            p = doc.add_paragraph()
            run(p, s["notes"], 9.5, False, rgb["muted"], italic=True)

    # ---- footer: title on the left, page x of y on the right
    ft = sec.footer.add_table(1, 2, width)
    no_borders(ft)
    lp, fp = ft.rows[0].cells[0].paragraphs[0], ft.rows[0].cells[1].paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(lp, (c.get("title") or "")[:90], 8.5, False, rgb["muted"])
    run(fp, "Seite " if not en else "Page ", 8.5, False, rgb["muted"])

    def field(p, code):
        r = run(p, "", 8.5, False, rgb["muted"])
        for kind, txt_ in (("begin", None), (None, code), ("separate", None), (None, "1"), ("end", None)):
            if kind:
                e = OxmlElement("w:fldChar")
                e.set(qn("w:fldCharType"), kind)
                r._r.append(e)
            elif txt_ == code:
                e = OxmlElement("w:instrText")
                e.set(qn("xml:space"), "preserve")
                e.text = f" {code} "
                r._r.append(e)
            else:
                e = OxmlElement("w:t")
                e.text = txt_
                r._r.append(e)
    field(fp, "PAGE")
    run(fp, " / ", 8.5, False, rgb["muted"])
    field(fp, "NUMPAGES")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _docx_table_designed(doc, tb, st, shade, borders, margins, no_borders, width):
    from docx.shared import Pt, RGBColor
    head, rows = tb.get("head") or [], tb.get("rows") or []
    ncol = max([len(head)] + [len(r) for r in rows] + [1])
    t = doc.add_table(rows=0, cols=ncol)
    no_borders(t)
    size = 10 if ncol <= 5 else 8.5
    line = ("single", 4, st["line"])
    for r, row in enumerate(([head] if head else []) + rows):
        is_head = bool(head) and r == 0
        cells = t.add_row().cells
        for i in range(ncol):
            cell = cells[i]
            cell.text = ""
            if is_head:
                shade(cell, st["accent"])
            elif (r % 2 == 0) == bool(head):
                shade(cell, st["zebra"])
            borders(cell, bottom=line)
            margins(cell, 70, 110)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            v = row[i] if i < len(row) else ""
            rr = p.add_run(v)
            rr.font.size = Pt(size)
            rr.bold = is_head or (i == 0 and ncol > 1)
            rr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if is_head else RGBColor.from_string(st["fg"])
            if not is_head and i > 0 and number(v, "de") is not None:
                p.alignment = 2
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _docx_table(doc, tb):
    head, rows = tb.get("head") or [], tb.get("rows") or []
    ncol = max([len(head)] + [len(r) for r in rows] + [1])
    t = doc.add_table(rows=0, cols=ncol)
    try:
        t.style = "Light Grid Accent 1"
    except (KeyError, ValueError):
        t.style = "Table Grid"
    if head:
        cells = t.add_row().cells
        for i, h in enumerate(head):
            cells[i].text = h
            for r in cells[i].paragraphs[0].runs:
                r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row[:ncol]):
            cells[i].text = v
    doc.add_paragraph("")


def _docx_letter(doc, c, Pt, Cm, center):
    L = c["letter"]
    if L.get("sender"):
        p = doc.add_paragraph(" · ".join(L["sender"]))
        p.runs[0].font.size = Pt(8)
        p.runs[0].underline = True
    for line in L.get("recipient") or []:
        p = doc.add_paragraph(line)
        p.paragraph_format.space_after = Pt(0)
    doc.add_paragraph("")
    p = doc.add_paragraph(L.get("place_date") or c.get("date") or "")
    p.alignment = 2   # right
    p = doc.add_paragraph()
    r = p.add_run(L.get("subject") or c.get("title") or "")
    r.bold = True
    if L.get("greeting"):
        doc.add_paragraph(L["greeting"])
    for t in L.get("paras") or []:
        doc.add_paragraph(t)
    if L.get("closing"):
        doc.add_paragraph(L["closing"])
    doc.add_paragraph("")
    if L.get("signature"):
        doc.add_paragraph(L["signature"])


# ---------------------------------------------------------------- PowerPoint
def pptx(c, style, pictures=None, base=None):
    if not base:
        return _pptx_designed(c, style, pictures)
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Emu, Inches, Pt
    prs = Presentation(base) if base else Presentation()
    if base:     # an own template: its masters and layouts stay, its slides go
        lst = prs.slides._sldIdLst
        for sid in list(lst):
            prs.part.drop_rel(sid.rId)
            lst.remove(sid)
    else:
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W, H = prs.slide_width, prs.slide_height
    acc, acc2, fg, bg, mut = (RGBColor.from_string(style[k]) for k in ("accent", "accent2", "fg", "bg", "muted"))

    def layout(i):
        n = len(prs.slide_layouts)
        return prs.slide_layouts[min(i, n - 1)]

    def text(slide, x, y, w, h, value, size, color, bold=False):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.text = value
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.size, r.font.bold, r.font.color.rgb = Pt(size), bold, color
                if not base:
                    r.font.name = style["font"]
        return tf

    def plain(slide):
        if base:
            return
        f = slide.background.fill
        f.solid()
        f.fore_color.rgb = bg

    # title slide
    if base:
        s = prs.slides.add_slide(layout(0))
        if s.shapes.title is not None:
            s.shapes.title.text = c.get("title") or ""
        sub = next((p for p in s.placeholders if p.placeholder_format.idx == 1), None)
        if sub is not None:
            sub.text = c.get("subtitle") or " · ".join(x for x in (c.get("author"), c.get("date")) if x)
    else:
        s = prs.slides.add_slide(layout(6))
        f = s.background.fill
        f.solid()
        f.fore_color.rgb = acc
        bar = s.shapes.add_shape(1, 0, H - Inches(0.35), W, Inches(0.35))
        bar.fill.solid()
        bar.fill.fore_color.rgb = acc2
        bar.line.fill.background()
        white = RGBColor(0xFF, 0xFF, 0xFF)
        text(s, Inches(0.9), Inches(2.4), W - Inches(1.8), Inches(1.6), c.get("title") or "", 44, white, True)
        sub = c.get("subtitle") or ""
        info = " · ".join(x for x in (c.get("author"), c.get("date")) if x)
        text(s, Inches(0.9), Inches(4.1), W - Inches(1.8), Inches(1.2), "\n".join(x for x in (sub, info) if x), 20, white)
    for n, sec in enumerate(_sections(c), 2):
        pics = pictures(sec) if pictures else []
        if base:
            s = prs.slides.add_slide(layout(1))
            if s.shapes.title is not None:
                s.shapes.title.text = sec.get("heading") or ""
            body = next((p for p in s.placeholders if p.placeholder_format.idx == 1), None)
            lines = (sec.get("paras") or []) + (sec.get("bullets") or [])
            if body is not None and lines and not (sec.get("table") or sec.get("chart")):
                body.text_frame.text = lines[0]
                for t in lines[1:]:
                    body.text_frame.add_paragraph().text = t
            elif body is not None:
                body._element.getparent().remove(body._element)
            top, left, width = Inches(1.6), Inches(0.6), W - Inches(1.2)
            if lines and (sec.get("table") or sec.get("chart")):
                text(s, left, top, width / 2 - Inches(0.2), H - top - Inches(0.6), "\n".join("• " + t for t in lines), 16, fg)
                left, width = left + width / 2, width / 2
        else:
            s = prs.slides.add_slide(layout(6))
            plain(s)
            line = s.shapes.add_shape(1, Inches(0.6), Inches(1.25), Inches(1.2), Inches(0.08))
            line.fill.solid()
            line.fill.fore_color.rgb = acc2
            line.line.fill.background()
            text(s, Inches(0.6), Inches(0.35), W - Inches(1.2), Inches(0.9), sec.get("heading") or "", 30, acc, True)
            num = text(s, W - Inches(1.1), H - Inches(0.55), Inches(0.8), Inches(0.4), str(n), 11, mut)
            num.paragraphs[0].alignment = 3
            lines = (sec.get("paras") or []) + ["• " + t for t in (sec.get("bullets") or [])]
            side = bool(sec.get("table") or sec.get("chart") or pics)
            top, left, width = Inches(1.6), Inches(0.6), W - Inches(1.2)
            if lines:
                bw = width * (0.45 if side else 1)
                tf = text(s, left, top, bw, H - top - Inches(0.8), "\n".join(lines), 20 if len(lines) <= 6 else 16, fg)
                for p in tf.paragraphs:
                    p.space_after = Pt(8)
                if side:
                    left, width = left + bw + Inches(0.3), width - bw - Inches(0.3)
        h = H - top - Inches(0.8)
        if sec.get("chart"):
            ch = sec["chart"]
            cd = CategoryChartData()
            cd.categories = ch.get("labels") or []
            for sr in (ch.get("series") or [])[:1 if ch.get("type") == "pie" else 6]:
                cd.add_series(sr.get("label") or "", [float(v) for v in sr.get("values") or []])
            kind = {"pie": XL_CHART_TYPE.PIE, "line": XL_CHART_TYPE.LINE_MARKERS}.get(ch.get("type"), XL_CHART_TYPE.COLUMN_CLUSTERED)
            gh = h if not sec.get("table") else h / 2
            gf = s.shapes.add_chart(kind, left, top, width, gh, cd)
            chart = gf.chart
            chart.has_legend = len(ch.get("series") or []) > 1 or ch.get("type") == "pie"
            if chart.has_legend:
                chart.legend.position = XL_LEGEND_POSITION.BOTTOM
                chart.legend.include_in_layout = False
            if ch.get("title"):
                chart.has_title = True
                chart.chart_title.text_frame.text = ch["title"]
            if not base:
                try:
                    chart.font.color.rgb = fg
                    chart.font.size = Pt(12)
                    if kind != XL_CHART_TYPE.PIE:
                        for i, pl in enumerate(chart.plots[0].series):
                            col = acc if i == 0 else acc2
                            if kind == XL_CHART_TYPE.COLUMN_CLUSTERED:
                                pl.format.fill.solid()
                                pl.format.fill.fore_color.rgb = col
                            else:
                                pl.format.line.color.rgb = col
                except (AttributeError, ValueError, IndexError):
                    pass
            if sec.get("table"):
                top, h = top + gh + Inches(0.2), h / 2 - Inches(0.2)
        if sec.get("table"):
            tb = sec["table"]
            head, rows = tb.get("head") or [], (tb.get("rows") or [])[:14]
            ncol = max([len(head)] + [len(r) for r in rows] + [1])
            nrow = len(rows) + (1 if head else 0)
            if nrow:
                gt = s.shapes.add_table(nrow, min(ncol, 12), left, top, width, min(h, Emu(Inches(0.42) * nrow))).table
                for r, row in enumerate(([head] if head else []) + rows):
                    for i in range(min(ncol, 12)):
                        cell = gt.cell(r, i)
                        cell.text = row[i] if i < len(row) else ""
                        for p in cell.text_frame.paragraphs:
                            for run in p.runs:
                                run.font.size = Pt(12 if nrow < 9 else 10)
        elif pics and not sec.get("chart"):
            k = len(pics)
            cols = 1 if k == 1 else 2 if k <= 4 else 3
            rows_ = math.ceil(k / cols)
            cw, chh = width / cols, h / rows_
            for i, p in enumerate(pics):
                try:
                    from PIL import Image
                    with Image.open(p) as im:
                        iw, ih = im.size
                except Exception:
                    continue
                scale = min((cw - Inches(0.15)) / iw, (chh - Inches(0.15)) / ih)
                pw, ph = int(iw * scale), int(ih * scale)
                x = left + cw * (i % cols) + (cw - pw) / 2
                y = top + chh * (i // cols) + (chh - ph) / 2
                s.shapes.add_picture(p, int(x), int(y), pw, ph)
        if sec.get("notes"):
            s.notes_slide.notes_text_frame.text = sec["notes"]
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def _est_lines(text_, size, width_in):
    """About how many lines a text needs in a box this wide (a rough measure, enough to pick sizes)."""
    per = max(8, int(width_in * 72 / (size * 0.5)))
    return sum(max(1, math.ceil(len(part) / per)) for part in str(text_ or "").split("\n"))


def _pptx_designed(c, style, pictures=None):
    """The Spark's own slides: a title slide with shapes, an overview, and per part a layout that fits what it holds
    (cards for a few points, big numbers, a core sentence in a box, chart and table in a card)."""
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Inches, Pt
    from lxml import etree

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W, H = prs.slide_width, prs.slide_height
    col = {k: RGBColor.from_string(style[k]) for k in ("accent", "accent2", "fg", "bg", "muted", "tint", "card", "line")}
    white = RGBColor(0xFF, 0xFF, 0xFF)
    en = c.get("lang") == "en"
    blank = prs.slide_layouts[6]
    IN = Inches(1)

    def new_slide(bg=None):
        sl = prs.slides.add_slide(blank)
        f = sl.background.fill
        f.solid()
        f.fore_color.rgb = bg or col["bg"]
        return sl

    def box(sl, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, line=None, radius=None, alpha=None):
        sh = sl.shapes.add_shape(shape, int(x), int(y), int(w), int(h))
        sh.shadow.inherit = False
        if fill is None:
            sh.fill.background()
        else:
            sh.fill.solid()
            sh.fill.fore_color.rgb = fill
            if alpha is not None:      # see-through: <a:alpha> inside the color
                clr = sh.fill._xPr.find(qn("a:solidFill"))[0]
                etree.SubElement(clr, qn("a:alpha")).set("val", str(int(alpha * 100000)))
        if line is None:
            sh.line.fill.background()
        else:
            sh.line.color.rgb = line
            sh.line.width = Pt(1)
        if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            sh.adjustments[0] = radius
        if sh.has_text_frame:
            sh.text_frame.text = ""
        return sh

    def text(sl, x, y, w, h, value, size, color, bold=False, head=False, align=None, anchor=None, spacing=None):
        tb = sl.shapes.add_textbox(int(x), int(y), int(w), int(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        if anchor:
            tf.vertical_anchor = anchor
        lines = str(value or "").split("\n")
        for k, ln in enumerate(lines):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            if align is not None:
                p.alignment = align
            if spacing:
                p.space_after = Pt(spacing)
            r = p.add_run()
            r.text = ln
            r.font.size, r.font.bold, r.font.color.rgb = Pt(size), bold, color
            r.font.name = style["head"] if head else style["font"]
        return tf

    def bullets(sl, x, y, w, h, items, size, color=None):
        """A list with round markers in the accent color (real PowerPoint bullets, so editing keeps them)."""
        tf = text(sl, x, y, w, h, "\n".join(items), size, color or col["fg"], spacing=size * 0.55)
        for p in tf.paragraphs:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(Pt(size * 1.25))))
            pPr.set("indent", str(-int(Pt(size * 1.25))))
            for tag in ("a:buClr", "a:buFont", "a:buChar"):
                old = pPr.find(qn(tag))
                if old is not None:
                    pPr.remove(old)
            bc = etree.SubElement(pPr, qn("a:buClr"))
            etree.SubElement(bc, qn("a:srgbClr")).set("val", style["accent"])
            etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
            etree.SubElement(pPr, qn("a:buChar")).set("char", "●")
        return tf

    def fit(items, w_in, h_in, sizes=(24, 22, 20, 18, 16, 14, 12)):
        """The biggest of sizes at which these lines fit the box."""
        longest = max([len(w_) for t in items for w_ in str(t).split()] + [1])
        for sz in sizes:
            need = sum(_est_lines(t, sz, w_in) for t in items) * sz * 1.25 / 72 + len(items) * sz * 0.55 / 72
            if need <= h_in and longest * sz * 0.52 / 72 <= w_in:     # no word broken in two
                return sz
        return sizes[-1]

    total = 1 + len(_sections(c))
    agenda = len(_sections(c)) >= 4
    total += 1 if agenda else 0
    title_ = c.get("title") or ""

    def chrome(sl, n, heading, kicker):
        """Heading with its small number above, the accent mark, the footer with title and page number."""
        box(sl, 0, 0, Inches(0.18), H, col["accent"])
        text(sl, Inches(0.7), Inches(0.42), Inches(6), Inches(0.35), kicker, 12, col["accent2"], bold=True)
        size = 32 if len(heading) <= 45 else 26 if len(heading) <= 80 else 22
        text(sl, Inches(0.7), Inches(0.72), W - Inches(1.4), Inches(1.0), heading, size, col["fg"], bold=True, head=True)
        box(sl, Inches(0.7), Inches(1.72), Inches(0.9), Inches(0.06), col["accent"])
        text(sl, Inches(0.7), H - Inches(0.5), W - Inches(2.5), Inches(0.3), title_, 10, col["muted"])
        text(sl, W - Inches(1.6), H - Inches(0.5), Inches(0.9), Inches(0.3), f"{n} / {total}", 10, col["muted"],
             align=PP_ALIGN.RIGHT)

    # ---- title slide: the accent color, two big see-through circles, title, subtitle, who and when
    dark = style["dark"]
    sl = new_slide(col["bg"] if dark else col["accent"])
    box(sl, W - Inches(5.2), -Inches(2.2), Inches(7.5), Inches(7.5), col["accent2"], MSO_SHAPE.OVAL, alpha=0.45 if dark else 0.55)
    box(sl, W - Inches(2.6), H - Inches(3.0), Inches(4.4), Inches(4.4), col["accent"] if dark else white, MSO_SHAPE.OVAL,
        alpha=0.35 if dark else 0.12)
    tw = col["fg"] if dark else white
    box(sl, Inches(0.9), Inches(2.05), Inches(1.1), Inches(0.08), col["accent"] if dark else tw)
    kind_name = {"praesentation": ("PRÄSENTATION", "PRESENTATION"), "bericht": ("BERICHT", "REPORT"),
                 "zusammenfassung": ("ZUSAMMENFASSUNG", "SUMMARY"), "sonstiges": ("", "")}.get(c.get("kind"), ("", ""))[1 if en else 0]
    if kind_name:
        text(sl, Inches(0.9), Inches(1.55), Inches(6), Inches(0.4), kind_name, 14, tw, bold=True)
    tsize = 54 if len(title_) <= 30 else 44 if len(title_) <= 60 else 34
    text(sl, Inches(0.9), Inches(2.35), Inches(8.6), Inches(2.6), title_, tsize, tw, bold=True, head=True,
         anchor=MSO_ANCHOR.TOP)
    sub = c.get("subtitle") or ""
    info = " · ".join(x for x in (c.get("author"), c.get("date")) if x)
    if sub:
        text(sl, Inches(0.9), Inches(5.05), Inches(8.6), Inches(0.8), sub, 24, tw)
    if info:
        text(sl, Inches(0.9), H - Inches(1.0), Inches(8.6), Inches(0.4), info, 14, tw)

    n = 1
    # ---- overview: the parts as numbered cards
    if agenda:
        n += 1
        sl = new_slide()
        chrome(sl, n, "Überblick" if not en else "Overview", f"{0:02d}")
        heads = [s_.get("heading") or "" for s_ in _sections(c)][:12]
        cols_ = 2 if len(heads) > 4 else 1
        per = math.ceil(len(heads) / cols_)
        top, avail = Inches(2.15), H - Inches(2.15) - Inches(0.85)
        rh = min(Inches(0.9), int(avail / per))
        cw = (W - Inches(1.4) - Inches(0.4) * (cols_ - 1)) / cols_
        for k, hd in enumerate(heads):
            x = Inches(0.7) + (cw + Inches(0.4)) * (k // per)
            y = top + rh * (k % per)
            box(sl, x, y + Inches(0.06), cw, rh - Inches(0.12), col["card"], MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.18)
            d = min(rh - Inches(0.3), Inches(0.55))
            num = box(sl, x + Inches(0.15), y + (rh - d) / 2, d, d, col["accent"], MSO_SHAPE.OVAL)
            tf = num.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(k + 1)
            r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(16), True, white, style["font"]
            text(sl, x + Inches(0.3) + d, y, cw - d - Inches(0.5), rh, hd, 18 if len(hd) < 50 else 15, col["fg"],
                 anchor=MSO_ANCHOR.MIDDLE)

    for k_sec, sec in enumerate(_sections(c), 1):
        n += 1
        sl = new_slide()
        heading = sec.get("heading") or ""
        chrome(sl, n, heading, f"{k_sec:02d}")
        pics = pictures(sec) if pictures else []
        paras, bl = sec.get("paras") or [], sec.get("bullets") or []
        stats, quote = sec.get("stats") or [], sec.get("quote") or ""
        visual = bool(sec.get("chart") or sec.get("table") or pics)
        L, R = Inches(0.7), W - Inches(0.7)
        top, bottom = Inches(2.1), H - Inches(0.75)
        width = R - L

        def lead(x, y, w, maxh):
            """The paragraphs as a calm lead text; returns the height used."""
            if not paras:
                return 0
            t = "\n".join(paras)
            sz = fit([t], w / IN, maxh / IN, (20, 18, 16, 14, 12))
            hh = min(maxh, int(Pt(sz) * 1.3 * sum(_est_lines(x_, sz, w / IN) for x_ in paras) + Pt(sz) * 0.6 * len(paras)))
            text(sl, x, y, w, hh, t, sz, col["muted"], spacing=sz * 0.5)
            return hh + Inches(0.2)

        def stat_cards(x, y, w, h):
            k = len(stats)
            gap = Inches(0.3)
            cw = (w - gap * (k - 1)) / k
            for i, st in enumerate(stats):
                cx = x + (cw + gap) * i
                box(sl, cx, y, cw, h, col["tint"], MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
                box(sl, cx, y, Inches(0.09), h, col["accent"] if i % 2 == 0 else col["accent2"])
                vsz = 48 if len(st["value"]) <= 7 and h > Inches(1.7) else 36 if len(st["value"]) <= 10 else 26
                text(sl, cx + Inches(0.4), y + Inches(0.25), cw - Inches(0.6), Inches(vsz * 1.25 / 72), st["value"], vsz,
                     col["accent"], bold=True, head=True)
                text(sl, cx + Inches(0.4), y + Inches(0.35 + vsz * 1.25 / 72), cw - Inches(0.6),
                     h - Inches(0.5 + vsz * 1.25 / 72), st["label"], 16, col["fg"])

        def quote_box(x, y, w, h):
            box(sl, x, y, w, h, col["tint"], MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
            box(sl, x, y, Inches(0.1), h, col["accent"])
            text(sl, x + Inches(0.35), y - Inches(0.25), Inches(1), Inches(1), "“", 72, col["accent"], bold=True,
                 head=True)
            sz = fit([quote], (w - Inches(1.4)) / IN, (h - Inches(0.3)) / IN, (28, 24, 22, 20, 18, 16))
            text(sl, x + Inches(1.1), y + Inches(0.15), w - Inches(1.4), h - Inches(0.3), quote, sz, col["fg"],
                 head=True, anchor=MSO_ANCHOR.MIDDLE)

        def cards(x, y, w, h, items):
            """A few points side by side (up to 4) or in two rows, each in its own card with a number."""
            k = len(items)
            per_row = k if k <= 4 else math.ceil(k / 2)
            rows = math.ceil(k / per_row)
            gap = Inches(0.3)
            cw = (w - gap * (per_row - 1)) / per_row
            ch = min((h - gap * (rows - 1)) / rows, Inches(3.2))
            sz = min(fit([t], (cw - Inches(0.6)) / IN, (ch - Inches(1.2)) / IN, (24, 22, 20, 18, 16, 14, 12)) for t in items)
            need = max(_est_lines(t, sz, (cw - Inches(0.6)) / IN) for t in items) * sz * 1.3 / 72
            ch = min(ch, max(Inches(1.35 + need), int(min(Inches(2.1), ch))))
            for i, t in enumerate(items):
                cx, cy = x + (cw + gap) * (i % per_row), y + (ch + gap) * (i // per_row)
                box(sl, cx, cy, cw, ch, col["card"], MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
                box(sl, cx, cy, cw, Inches(0.08), col["accent"] if i % 2 == 0 else col["accent2"])
                d = Inches(0.5)
                num = box(sl, cx + Inches(0.3), cy + Inches(0.35), d, d, col["accent"] if i % 2 == 0 else col["accent2"],
                          MSO_SHAPE.OVAL)
                tf = num.text_frame
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER
                r = p.add_run()
                r.text = str(i + 1)
                r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(16), True, white, style["font"]
                text(sl, cx + Inches(0.3), cy + Inches(1.05), cw - Inches(0.6), ch - Inches(1.2), t, sz, col["fg"])
            return ch * rows + gap * (rows - 1)

        y = top
        if visual:
            # text on the left (lead, numbers, points), the chart, table or pictures in a card on the right
            textual = bool(paras or bl or stats or quote)
            lw = width * 0.38 if textual else 0
            vx = L + (lw + Inches(0.4) if textual else 0)
            vw = R - vx
            if textual:
                ty = top
                ty += lead(L, ty, lw, Inches(1.6))
                if stats:
                    sh_ = Inches(1.3)
                    for i, st in enumerate(stats[:3]):
                        box(sl, L, ty, lw, sh_ - Inches(0.15), col["tint"], MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
                        text(sl, L + Inches(0.25), ty + Inches(0.1), lw - Inches(0.4), Inches(0.6), st["value"], 28,
                             col["accent"], bold=True, head=True)
                        text(sl, L + Inches(0.25), ty + Inches(0.7), lw - Inches(0.4), Inches(0.45), st["label"], 13, col["fg"])
                        ty += sh_
                if bl:
                    avail = bottom - ty - (Inches(1.4) if quote else 0)
                    sz = fit(bl, lw / IN - 0.3, avail / IN, (20, 18, 16, 14, 12))
                    bullets(sl, L, ty, lw, avail, bl, sz)
                    ty += avail
                if quote:
                    quote_box(L, bottom - Inches(1.3), lw, Inches(1.3))
            vy, vh = top, bottom - top
            if sec.get("table") and not sec.get("chart"):     # the card as high as the table
                tb_ = sec["table"]
                nr = min(14, len(tb_.get("rows") or [])) + (1 if tb_.get("head") else 0)
                nc = max([len(tb_.get("head") or [])] + [len(r) for r in tb_.get("rows") or []] + [1])
                tsz = 16 if nr <= 5 and nc <= 4 else 14 if nr <= 8 else 11
                vh = min(vh, int(Pt(tsz) * 2.2 * nr + Inches(0.5)))
            box(sl, vx, vy, vw, vh, col["card"], MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
            ix, iy, iw, ih = vx + Inches(0.3), vy + Inches(0.25), vw - Inches(0.6), vh - Inches(0.5)
            if sec.get("chart") and sec.get("table"):
                _pptx_chart(sl, sec["chart"], ix, iy, iw, ih * 0.55, style, col, CategoryChartData, XL_CHART_TYPE,
                            XL_LEGEND_POSITION, XL_LABEL_POSITION, Pt)
                _pptx_table(sl, sec["table"], ix, iy + ih * 0.6, iw, ih * 0.4, style, col, white, Pt, Emu, Inches)
            elif sec.get("chart"):
                _pptx_chart(sl, sec["chart"], ix, iy, iw, ih, style, col, CategoryChartData, XL_CHART_TYPE,
                            XL_LEGEND_POSITION, XL_LABEL_POSITION, Pt)
            elif sec.get("table"):
                _pptx_table(sl, sec["table"], ix, iy, iw, ih, style, col, white, Pt, Emu, Inches)
            else:
                _pptx_pictures(sl, pics, ix, iy, iw, ih, Inches)
        else:
            if stats and not bl and not paras and not quote:
                hh = min(Inches(2.6), bottom - top - Inches(0.8))
                stat_cards(L, top + (bottom - top - hh) * 0.35, width, hh)
            elif not bl and not stats and (quote or paras) and len(" ".join(paras + [quote])) <= 320:
                # one core sentence or a short text: a big statement in a box
                if not quote:
                    quote = " ".join(paras)
                    paras = []
                hh = min(Inches(3.2), bottom - top - Inches(0.6))
                quote_box(L, top + (bottom - top - hh) * 0.35, width, hh)
            else:
                y += lead(L, y, width, Inches(1.8))
                if stats:
                    sh_ = Inches(1.9) if not bl else Inches(1.6)
                    stat_cards(L, y, width, sh_)
                    y += sh_ + Inches(0.3)
                qh = Inches(1.2) if quote else 0
                avail = bottom - y - (qh + Inches(0.25) if quote else 0)
                if bl:
                    short = all(len(t) <= 110 for t in bl)
                    if 2 <= len(bl) <= 6 and short and avail >= Inches(2.2):
                        y += cards(L, y, width, avail, bl) + Inches(0.35)
                    else:
                        sz = fit(bl, width / IN - 0.4, avail / IN)
                        bullets(sl, L, y, width, avail, bl, sz)
                        y = bottom - qh
                if quote:
                    quote_box(L, min(y, bottom - qh), width, qh)
        if sec.get("notes"):
            sl.notes_slide.notes_text_frame.text = sec["notes"]
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def _pptx_chart(sl, ch, x, y, w, h, style, col, CategoryChartData, XL_CHART_TYPE, XL_LEGEND_POSITION,
                XL_LABEL_POSITION, Pt):
    from pptx.dml.color import RGBColor
    cd = CategoryChartData()
    cd.categories = ch.get("labels") or []
    pie = ch.get("type") == "pie"
    series = (ch.get("series") or [])[:1 if pie else 6]
    for sr in series:
        cd.add_series(sr.get("label") or "", [float(v) for v in sr.get("values") or []])
    kind = {"pie": XL_CHART_TYPE.DOUGHNUT, "line": XL_CHART_TYPE.LINE_MARKERS}.get(ch.get("type"),
                                                                                   XL_CHART_TYPE.COLUMN_CLUSTERED)
    chart = sl.shapes.add_chart(kind, int(x), int(y), int(w), int(h), cd).chart
    palette = [RGBColor.from_string(h_) for h_ in chart_colors(style)]
    chart.font.size = Pt(12)
    chart.font.color.rgb = col["fg"]
    chart.font.name = style["font"]
    chart.has_legend = len(series) > 1 or pie
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(12)
    if ch.get("title"):
        chart.has_title = True
        tf = chart.chart_title.text_frame
        tf.text = ch["title"]
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.size, r.font.bold, r.font.color.rgb = Pt(16), True, col["fg"]
    else:
        chart.has_title = False
    try:
        plot = chart.plots[0]
        if pie:
            for i, pt in enumerate(plot.series[0].points):
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = palette[i % len(palette)]
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.show_percentage, dl.show_value, dl.number_format, dl.number_format_is_linked = True, False, "0%", False
            dl.font.size, dl.font.bold = Pt(12), True
        else:
            for i, sr in enumerate(plot.series):
                if kind == XL_CHART_TYPE.COLUMN_CLUSTERED:
                    sr.format.fill.solid()
                    sr.format.fill.fore_color.rgb = palette[i % len(palette)]
                else:
                    sr.format.line.color.rgb = palette[i % len(palette)]
                    sr.format.line.width = Pt(3.5)
                    sr.smooth = False
            if kind == XL_CHART_TYPE.COLUMN_CLUSTERED:
                plot.gap_width = 60
                if len(series) == 1 and len(ch.get("labels") or []) <= 12:
                    plot.has_data_labels = True
                    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
                    plot.data_labels.font.size = Pt(12)
                    plot.data_labels.font.bold = True
            va = chart.value_axis
            va.has_major_gridlines = True
            va.major_gridlines.format.line.color.rgb = col["line"]
            va.format.line.fill.background()
            va.tick_labels.font.color.rgb = col["muted"]
            ca = chart.category_axis
            ca.format.line.color.rgb = col["line"]
            ca.tick_labels.font.color.rgb = col["muted"]
    except (AttributeError, ValueError, IndexError, KeyError):
        pass


def _pptx_table(sl, tb, x, y, w, h, style, col, white, Pt, Emu, Inches):
    head, rows = tb.get("head") or [], (tb.get("rows") or [])[:14]
    ncol = min(12, max([len(head)] + [len(r) for r in rows] + [1]))
    nrow = len(rows) + (1 if head else 0)
    if not nrow:
        return
    size = 16 if nrow <= 5 and ncol <= 4 else 14 if nrow <= 8 else 11
    rh = Emu(int(Pt(size) * 2.2))
    shape = sl.shapes.add_table(nrow, ncol, int(x), int(y), int(w), int(min(h, rh * nrow)))
    t = shape.table
    t.first_row = bool(head)
    t.horz_banding = False
    from pptx.oxml.ns import qn
    from lxml import etree
    tblPr = t._tbl.tblPr      # "No Style, No Grid": our own fills, no white lines on a dark slide
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    for r, row in enumerate(([head] if head else []) + rows):
        is_head = bool(head) and r == 0
        for i in range(ncol):
            cell = t.cell(r, i)
            cell.text = row[i] if i < len(row) else ""
            cell.fill.solid()
            cell.fill.fore_color.rgb = col["accent"] if is_head else (col["bg"] if r % 2 else col["tint"]) \
                if head else (col["tint"] if r % 2 else col["bg"])
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(size)
                    run.font.name = style["font"]
                    run.font.bold = is_head or (i == 0 and not is_head)
                    run.font.color.rgb = white if is_head else col["fg"]


def _pptx_pictures(sl, pics, left, top, width, h, Inches):
    k = len(pics)
    if not k:
        return
    cols = 1 if k == 1 else 2 if k <= 4 else 3
    rows_ = math.ceil(k / cols)
    cw, chh = width / cols, h / rows_
    for i, p in enumerate(pics):
        try:
            from PIL import Image
            with Image.open(p) as im:
                iw, ih = im.size
        except Exception:
            continue
        scale = min((cw - Inches(0.15)) / iw, (chh - Inches(0.15)) / ih)
        pw, ph = int(iw * scale), int(ih * scale)
        x = left + cw * (i % cols) + (cw - pw) / 2
        y = top + chh * (i // cols) + (chh - ph) / 2
        sl.shapes.add_picture(p, int(x), int(y), pw, ph)


# ---------------------------------------------------------------- Excel
def xlsx(c, style, pictures=None, base=None):
    import xlsxwriter
    style = _print_style(style)
    out = io.BytesIO()
    lang = c.get("lang") or "de"
    wb = xlsxwriter.Workbook(out, {"in_memory": True, "strings_to_formulas": False, "strings_to_urls": False,
                                   "strings_to_numbers": False})
    wb.set_properties({"title": (c.get("title") or "")[:200], "author": (c.get("author") or "")[:80]})
    fn, acc, line = style["font"], "#" + style["accent"], "#" + style["line"]
    base_f = {"font_name": fn, "font_size": 10.5, "font_color": "#" + style["fg"], "valign": "vcenter", "indent": 1}

    def fmt(**kw):
        return wb.add_format({**base_f, **kw})
    head = fmt(bold=True, font_color="#FFFFFF", bg_color=acc, bottom=2, bottom_color=acc, text_wrap=True)
    headn = fmt(bold=True, font_color="#FFFFFF", bg_color=acc, bottom=2, bottom_color=acc, align="right")
    cells, nums = {}, {}
    for zebra in (False, True):          # body cells: thin line below, every second row tinted
        bg = {"bg_color": "#" + style["zebra"]} if zebra else {}
        for total in (False, True):
            edge = {"top": 2, "top_color": acc, "bold": True} if total else {}
            cells[zebra, total] = fmt(bottom=1, bottom_color=line, text_wrap=True, **bg, **edge)
            nums[zebra, total] = fmt(bottom=1, bottom_color=line, num_format="#,##0.##", align="right", **bg, **edge)
    titlef = fmt(bold=True, font_size=18, font_color="#FFFFFF", bg_color=acc, font_name=style["head"], indent=1)
    subf = fmt(font_size=10, font_color="#" + mix(style["accent"], "FFFFFF", 0.8), bg_color=acc, indent=1)
    textf = fmt(text_wrap=True, valign="top")
    boldf = fmt(bold=True, font_size=13, font_color=acc, font_name=style["head"], bottom=2, bottom_color=acc)
    TOTAL = re.compile(r"^(summe|gesamt|insgesamt|total|sum|ergebnis)\b", re.I)

    def banner(ws, text, sub, cols):
        """A colored title band over the sheet's width, gridlines off."""
        ws.hide_gridlines(2)
        ws.set_tab_color(acc)
        ws.set_paper(9)
        ws.fit_to_pages(1, 0)
        ws.set_margins(0.5, 0.5, 0.6, 0.6)
        ws.set_row(0, 34)
        ws.merge_range(0, 0, 0, max(cols - 1, 0), text, titlef) if cols > 1 else ws.write_string(0, 0, text, titlef)
        ws.set_row(1, 18)
        ws.merge_range(1, 0, 1, max(cols - 1, 0), sub, subf) if cols > 1 else ws.write_string(1, 0, sub, subf)
        return 3
    used = set()

    def sheet_name(v, i):
        v = re.sub(r"[\[\]:*?/\\']", " ", str(v or "")).strip()[:28] or f"Blatt {i}"
        base_, k = v, 2
        while v.lower() in used:
            v = f"{base_[:25]} {k}"
            k += 1
        used.add(v.lower())
        return v

    over, orow = None, 0      # the sheet for sections without a table or chart, and its next row
    for i, s in enumerate(_sections(c), 1):
        tb, ch = s.get("table"), s.get("chart")
        if not tb and not ch:
            if over is None:
                over = wb.add_worksheet(sheet_name(c.get("title") or "Übersicht", 0))
                over.set_column(0, 0, 100)
                orow = banner(over, c.get("title") or "", " · ".join(
                    x for x in (c.get("subtitle"), c.get("author"), c.get("date")) if x), 1)
            over.write_string(orow, 0, s.get("heading") or "", boldf)
            orow += 1
            for t in (s.get("paras") or []) + ["• " + b for b in (s.get("bullets") or [])]:
                over.write_string(orow, 0, t, textf)
                orow += 1
            orow += 1
            continue
        ws = wb.add_worksheet(sheet_name(s.get("heading"), i))
        if not tb and ch:      # the chart's numbers become the table it is drawn from
            tb = {"head": [""] + [sr.get("label") or "" for sr in ch.get("series") or []],
                  "rows": [[lab] + [_fmt(sr["values"][k]) if k < len(sr.get("values") or []) else "" for sr in ch["series"]]
                           for k, lab in enumerate(ch.get("labels") or [])]}
        ncol = max([len(tb.get("head") or [])] + [len(r) for r in tb.get("rows") or []] + [1])
        r0 = banner(ws, s.get("heading") or c.get("title") or "", " · ".join(
            x for x in ((c.get("title") if s.get("heading") else c.get("subtitle")), c.get("date")) if x), ncol)
        for t in (s.get("paras") or []) + ["• " + b for b in (s.get("bullets") or [])]:
            ws.write_string(r0, 0, t, textf)
            r0 += 1
        if s.get("paras") or s.get("bullets"):
            r0 += 1
        hd, rows = tb.get("head") or [], tb.get("rows") or []
        widths = [10] * ncol
        numeric = [k > 0 and bool(rows) and all(number(r[k], lang) is not None or formula(r[k]) or not r[k]
                                                for r in rows if k < len(r)) for k in range(ncol)]
        if hd:
            ws.set_row(r0, 22)
            for k in range(ncol):
                v = hd[k] if k < len(hd) else ""
                ws.write_string(r0, k, v, headn if numeric[k] else head)
                widths[k] = max(widths[k], len(v) + 4)
        first = r0 + (1 if hd else 0)
        for ri, row in enumerate(rows):
            ws.set_row(first + ri, 20)
            key = (ri % 2 == 1, bool(row) and bool(TOTAL.match(str(row[0]).strip())))
            for k in range(ncol):
                v = row[k] if k < len(row) else ""
                f, n = formula(v), number(v, lang)
                if f:
                    ws.write_formula(first + ri, k, f, nums[key])
                elif n is not None:
                    ws.write_number(first + ri, k, n, nums[key])
                else:
                    ws.write_string(first + ri, k, v, cells[key])
                widths[k] = max(widths[k], min(50, len(v) + 4))
        for k, w in enumerate(widths):
            ws.set_column(k, k, w)
        if hd:
            ws.freeze_panes(first, 0)
            if rows:
                ws.autofilter(r0, 0, first + len(rows) - 1, ncol - 1)
        if ch and rows:
            kind = {"pie": "pie", "line": "line"}.get(ch.get("type"), "column")
            chart = wb.add_chart({"type": kind})
            series = [k for k in range(1, ncol) if numeric[k]][:1 if kind == "pie" else 6] or [1]
            cols_ = ["#" + x for x in chart_colors(style)]
            last = first + len(rows) - 1
            if TOTAL.match(str(rows[-1][0] if rows[-1] else "").strip()) and len(rows) > 1:
                last -= 1      # a total row is not a data point
            font = {"name": fn, "color": "#" + style["muted"], "size": 9}
            nser = len(series)
            for j_, k in enumerate(series):
                col = cols_[j_ % len(cols_)]
                ser = {"name": [ws.name, r0, k] if hd else f"Reihe {k}",
                       "categories": [ws.name, first, 0, last, 0],
                       "values": [ws.name, first, k, last, k],
                       "data_labels": {"value": kind != "pie", "percentage": kind == "pie",
                                       "font": {"name": fn, "size": 8, "color": "#" + style["fg"]}}}
                if kind == "column":
                    ser.update(fill={"color": col}, border={"none": True}, gap=60)
                elif kind == "line":
                    ser.update(line={"color": col, "width": 2.5},
                               marker={"type": "circle", "size": 7, "fill": {"color": "#FFFFFF"},
                                       "border": {"color": col, "width": 2}})
                else:
                    ser["points"] = [{"fill": {"color": cols_[j % len(cols_)]}, "border": {"color": "#FFFFFF"}}
                                     for j in range(last - first + 1)]
                chart.add_series(ser)
            if kind != "pie":
                chart.set_x_axis({"num_font": font, "line": {"color": line}, "major_tick_mark": "none"})
                chart.set_y_axis({"num_font": font, "line": {"none": True}, "major_tick_mark": "none",
                                  "major_gridlines": {"visible": True, "line": {"color": line, "width": 0.75}}})
            chart.set_legend({"none": True} if nser == 1 and kind != "pie" else
                             {"position": "bottom", "font": font})
            chart.set_chartarea({"border": {"none": True}, "fill": {"color": "#FFFFFF"}})
            chart.set_plotarea({"border": {"none": True}})
            if ch.get("title"):
                chart.set_title({"name": ch["title"], "name_font": {"name": style["head"], "size": 13, "bold": True,
                                                                     "color": "#" + style["fg"]}})
            else:
                chart.set_title({"none": True})
            chart.set_size({"width": 640, "height": 360})
            ws.insert_chart(first + len(rows) + 1, 0, chart, {"x_offset": 4, "y_offset": 8})
    if not wb.worksheets():
        ws = wb.add_worksheet("Blatt 1")
        banner(ws, c.get("title") or "", "", 1)
    wb.close()
    return out.getvalue()


# ---------------------------------------------------------------- PDF
def pdf(c, style, pictures=None, base=None):
    if not c.get("letter") and base is None:
        return _pdf_designed(c, style, pictures)
    from fpdf import FPDF
    from fpdf.fonts import FontFace
    reg, bold = _font_files()
    fam = "DejaVu" if reg else "Helvetica"

    def tx(s):
        s = str(s or "")
        if reg:
            return s
        return "".join(LATIN1.get(ch, ch) if ord(ch) > 255 or ch in LATIN1 else ch for ch in s).encode(
            "latin-1", "replace").decode("latin-1")

    doc = FPDF(format="A4")
    doc.set_auto_page_break(True, margin=18)
    doc.set_margins(22, 18, 18)
    if reg:
        doc.add_font("DejaVu", "", reg)
        doc.add_font("DejaVu", "B", bold)
    doc.set_title(tx(c.get("title"))[:200])
    if c.get("author"):
        doc.set_author(tx(c["author"])[:80])
    doc.set_creator("Spark")
    acc, fg, mut = _rgb(style["accent"] if not style["dark"] else "2F3A4F"), _rgb("1D2433"), _rgb("5F6B85")
    doc.add_page()
    W = doc.epw

    def para(t, size=11, b=False, color=fg, gap=2.5, align="L"):
        doc.set_font(fam, "B" if b else "", size)
        doc.set_text_color(*color)
        doc.multi_cell(W, size * 0.5, tx(t), align=align, new_x="LMARGIN", new_y="NEXT")
        doc.ln(gap)

    if c.get("letter"):
        L = c["letter"]
        if L.get("sender"):
            para(" · ".join(L["sender"]), 7.5, color=mut, gap=1)
        for line in L.get("recipient") or []:
            para(line, 11, gap=0)
        doc.ln(14)
        para(L.get("place_date") or c.get("date") or "", 11, align="R", gap=6)
        para(L.get("subject") or c.get("title") or "", 11, b=True, gap=5)
        if L.get("greeting"):
            para(L["greeting"])
        for t in L.get("paras") or []:
            para(t)
        if L.get("closing"):
            para(L["closing"], gap=12)
        if L.get("signature"):
            para(L["signature"])
    else:
        para(c.get("title") or "", 22, b=True, color=acc, gap=2)
        if c.get("subtitle"):
            para(c["subtitle"], 13, color=mut, gap=1)
        if c.get("date") or c.get("author"):
            para(" · ".join(x for x in (c.get("author"), c.get("date")) if x), 9, color=mut, gap=4)
        for s in _sections(c):
            if s.get("heading"):
                doc.ln(2)
                para(s["heading"], 14, b=True, color=acc, gap=1.5)
            for t in s.get("paras") or []:
                para(t)
            for t in s.get("bullets") or []:
                doc.set_font(fam, "", 11)
                doc.set_text_color(*fg)
                doc.cell(6, 5.5, "-" if not reg else "•")
                doc.multi_cell(W - 6, 5.5, tx(t), new_x="LMARGIN", new_y="NEXT")
                doc.ln(1)
            if s.get("table"):
                tb = s["table"]
                hd, rows = tb.get("head") or [], tb.get("rows") or []
                ncol = max([len(hd)] + [len(r) for r in rows] + [1])
                doc.set_font(fam, "", 9 if ncol > 4 else 10)
                doc.set_text_color(*fg)
                with doc.table(headings_style=FontFace(emphasis="BOLD", color=(255, 255, 255), fill_color=acc),
                               first_row_as_headings=bool(hd), line_height=5) as t:
                    for row in ([hd] if hd else []) + rows:
                        rr = t.row()
                        for i in range(ncol):
                            rr.cell(tx(row[i] if i < len(row) else ""))
                doc.ln(3)
            if s.get("chart"):
                doc.image(io.BytesIO(chart_png(s["chart"], dict(style, dark=False, bg="FFFFFF"))), w=min(W, 160))
                doc.ln(3)
            for p in (pictures(s) if pictures else []):
                try:
                    doc.image(p, w=min(W, 120))
                    doc.ln(3)
                except Exception:
                    continue
            if s.get("notes"):
                para(s["notes"], 9, color=mut)
    return bytes(doc.output())


def _pdf_designed(c, style, pictures=None):
    """The Spark's own PDF: the same layout as the Word file (title block, numbered headings, cards, box, tables
    with a colored head), page numbers in the footer."""
    from fpdf import FPDF
    from fpdf.fonts import FontFace
    st = _print_style(style)
    reg, bold = _font_files()
    fam = "DejaVu" if reg else "Helvetica"
    en = c.get("lang") == "en"

    def tx(s_):
        s_ = str(s_ or "")
        if reg:
            return s_
        return "".join(LATIN1.get(ch, ch) if ord(ch) > 255 or ch in LATIN1 else ch for ch in s_).encode(
            "latin-1", "replace").decode("latin-1")

    acc, acc2, fg, mut = (_rgb(st[k]) for k in ("accent", "accent2", "fg", "muted"))
    tint, zebra, line = _rgb(st["tint"]), _rgb(st["zebra"]), _rgb(st["line"])
    title_ = tx(c.get("title"))

    class Doc(FPDF):
        def footer(self):
            self.set_y(-13)
            self.set_font(fam, "", 8)
            self.set_text_color(*mut)
            self.cell(self.epw / 2, 5, title_[:90])
            self.cell(self.epw / 2, 5, f"{'Seite' if not en else 'Page'} {self.page_no()} / {{nb}}", align="R")

    doc = Doc(format="A4")
    doc.alias_nb_pages()
    doc.set_auto_page_break(True, margin=20)
    doc.set_margins(22, 18, 20)
    if reg:
        doc.add_font("DejaVu", "", reg)
        doc.add_font("DejaVu", "B", bold)
    doc.set_title(title_[:200])
    if c.get("author"):
        doc.set_author(tx(c["author"])[:80])
    doc.set_creator("Spark")
    doc.add_page()
    W, L = doc.epw, doc.l_margin

    def para(t, size=10.5, b=False, color=fg, gap=2.2, align="L", w=None, x=None):
        doc.set_font(fam, "B" if b else "", size)
        doc.set_text_color(*color)
        if x is not None:
            doc.set_x(x)
        doc.multi_cell(w or W, size * 0.5, tx(t), align=align, new_x="LMARGIN", new_y="NEXT")
        doc.ln(gap)

    def need(h):
        if doc.get_y() + h > doc.h - doc.b_margin:
            doc.add_page()

    def height_of(t, size, w):
        doc.set_font(fam, "", size)
        return len(doc.multi_cell(w, size * 0.5, tx(t), dry_run=True, output="LINES")) * size * 0.5

    # ---- title block
    label = KIND_LABEL.get(c.get("kind"), ("", ""))[1 if en else 0]
    tsize = 24 if len(title_) <= 50 else 19
    th = height_of(title_, tsize, W - 16)
    sub = tx(c.get("subtitle"))
    info = " · ".join(tx(x) for x in (c.get("author"), c.get("date")) if x)
    bh = 12 + (6 if label else 0) + th + (8 if sub else 0) + (6 if info else 0) + 4
    doc.set_fill_color(*acc)
    doc.rect(L, doc.get_y(), W, bh, style="F", round_corners=True, corner_radius=3)
    y = doc.get_y() + 9
    doc.set_xy(L + 8, y)
    if label:
        para(label, 8, True, _rgb(mix(st["accent"], "FFFFFF", 0.7)), 1, w=W - 16, x=L + 8)
    para(title_, tsize, True, (255, 255, 255), 1.5, w=W - 16, x=L + 8)
    if sub:
        para(sub, 12, False, (255, 255, 255), 1, w=W - 16, x=L + 8)
    if info:
        para(info, 8.5, False, _rgb(mix(st["accent"], "FFFFFF", 0.75)), 0, w=W - 16, x=L + 8)
    doc.set_y(y - 9 + bh + 8)

    for k, s in enumerate(_sections(c), 1):
        if s.get("heading"):
            need(22)
            doc.ln(2)
            doc.set_font(fam, "B", 14)
            doc.set_text_color(*acc2)
            num = f"{k:02d}"
            doc.cell(doc.get_string_width(num) + 3, 8, num)
            doc.set_text_color(*fg)
            doc.multi_cell(W - doc.get_string_width(num) - 3, 8, tx(s["heading"]), new_x="LMARGIN", new_y="NEXT")
            doc.set_draw_color(*acc)
            doc.set_line_width(0.5)
            doc.line(L, doc.get_y() + 0.5, L + W, doc.get_y() + 0.5)
            doc.ln(4)
        for t in s.get("paras") or []:
            para(t)
        if s.get("stats"):
            n = len(s["stats"])
            gap = 4
            cw = (W - gap * (n - 1)) / n
            hh = 24 if all(len(x.get("label") or "") <= 28 for x in s["stats"]) else 28
            need(hh + 4)
            y = doc.get_y()
            for i, x in enumerate(s["stats"]):
                cx = L + (cw + gap) * i
                doc.set_fill_color(*tint)
                doc.rect(cx, y, cw, hh, style="F", round_corners=True, corner_radius=2)
                doc.set_fill_color(*(acc if i % 2 == 0 else acc2))
                doc.rect(cx, y, 1.4, hh, style="F")
                doc.set_xy(cx + 5, y + 3.5)
                doc.set_font(fam, "B", 17)
                doc.set_text_color(*acc)
                doc.cell(cw - 8, 8, tx(x["value"]))
                doc.set_xy(cx + 5, y + 12.5)
                doc.set_font(fam, "", 8.5)
                doc.set_text_color(*fg)
                doc.multi_cell(cw - 8, 4, tx(x.get("label")), align="L")
            doc.set_y(y + hh + 5)
        for t in s.get("bullets") or []:
            need(8)
            y = doc.get_y()
            doc.set_fill_color(*acc)
            doc.ellipse(L + 1, y + 1.6, 2.2, 2.2, style="F")
            doc.set_xy(L + 6, y)
            para(t, gap=1.2, w=W - 6, x=L + 6)
        if s.get("quote"):
            q = ("„" if not en else "“") + tx(s["quote"]) + ("“" if not en else "”")
            qh = height_of(q, 12, W - 22) + 10
            need(qh + 4)
            doc.ln(1)
            y = doc.get_y()
            doc.set_fill_color(*tint)
            doc.rect(L, y, W, qh, style="F", round_corners=True, corner_radius=2)
            doc.set_fill_color(*acc)
            doc.rect(L, y, 1.8, qh, style="F")
            doc.set_xy(L + 10, y + 5)
            para(q, 12, True, fg, 0, w=W - 22, x=L + 10)
            doc.set_y(y + qh + 5)
        if s.get("table"):
            tb = s["table"]
            hd, rows = tb.get("head") or [], tb.get("rows") or []
            ncol = max([len(hd)] + [len(r) for r in rows] + [1])
            doc.set_font(fam, "", 8.5 if ncol > 4 else 9.5)
            doc.set_text_color(*fg)
            doc.set_draw_color(*line)
            doc.set_line_width(0.2)
            with doc.table(headings_style=FontFace(emphasis="BOLD", color=(255, 255, 255), fill_color=acc),
                           first_row_as_headings=bool(hd), line_height=5.5, borders_layout="HORIZONTAL_LINES",
                           padding=(1.5, 2.5)) as t:
                for r, row in enumerate(([hd] if hd else []) + rows):
                    body = r - (1 if hd else 0)
                    rr = t.row() if body < 0 else t.row(style=FontFace(
                        color=fg, fill_color=zebra if body % 2 else (255, 255, 255)))
                    for i in range(ncol):
                        rr.cell(tx(row[i] if i < len(row) else ""))
            doc.ln(4)
        if s.get("chart"):
            need(90)
            doc.image(io.BytesIO(chart_png(s["chart"], st)), w=min(W, 160), x=L + (W - min(W, 160)) / 2)
            doc.ln(4)
        for p in (pictures(s) if pictures else []):
            try:
                doc.image(p, w=min(W, 120), x=L + (W - min(W, 120)) / 2)
                doc.ln(3)
            except Exception:
                continue
        if s.get("notes"):
            para(s["notes"], 9, color=mut)
    return bytes(doc.output())


# ---------------------------------------------------------------- Markdown, text, CSV
def _md_cell(v):
    return str(v).replace("|", "\\|").replace("\n", " ")


def md(c, style=None, pictures=None, base=None, plain=False):
    out = []
    h1, h2, b = ("", "", "- ") if plain else ("# ", "## ", "- ")
    if c.get("letter"):
        L = c["letter"]
        out += [" · ".join(L.get("sender") or []), ""] if L.get("sender") else []
        out += list(L.get("recipient") or []) + ["", L.get("place_date") or c.get("date") or "", ""]
        out += [(L.get("subject") if plain else "**" + (L.get("subject") or "") + "**"), ""]
        out += ([L["greeting"], ""] if L.get("greeting") else [])
        for t in L.get("paras") or []:
            out += [t, ""]
        out += ([L["closing"], ""] if L.get("closing") else []) + ([L["signature"]] if L.get("signature") else [])
        return ("\n".join(out).strip() + "\n").encode("utf-8")
    out += [h1 + (c.get("title") or ""), ""]
    if c.get("subtitle"):
        out += [c["subtitle"] if plain else "_" + c["subtitle"] + "_", ""]
    if c.get("date") or c.get("author"):
        out += [" · ".join(x for x in (c.get("author"), c.get("date")) if x), ""]
    for s in _sections(c):
        if s.get("heading"):
            out += [h2 + s["heading"], ""]
        for t in s.get("paras") or []:
            out += [t, ""]
        if s.get("stats"):
            out += [b + (x["value"] if plain else "**" + x["value"] + "**") + (" – " + x["label"] if x.get("label") else "")
                    for x in s["stats"]] + [""]
        if s.get("quote"):
            out += [("» " + s["quote"] + " «") if plain else "> **" + s["quote"] + "**", ""]
        if s.get("bullets"):
            out += [b + t for t in s["bullets"]] + [""]
        tb = s.get("table")
        if not tb and s.get("chart"):
            ch = s["chart"]
            tb = {"head": [""] + [sr.get("label") or "" for sr in ch.get("series") or []],
                  "rows": [[lab] + [_fmt(sr["values"][k]) if k < len(sr.get("values") or []) else "" for sr in ch["series"]]
                           for k, lab in enumerate(ch.get("labels") or [])]}
            if ch.get("title"):
                out += [ch["title"] + ":", ""]
        if tb:
            hd, rows = tb.get("head") or [], tb.get("rows") or []
            ncol = max([len(hd)] + [len(r) for r in rows] + [1])
            if plain:
                for row in ([hd] if hd else []) + rows:
                    out.append("\t".join(row))
            else:
                hd = hd or [""] * ncol
                out.append("| " + " | ".join(_md_cell(x) for x in hd + [""] * (ncol - len(hd))) + " |")
                out.append("|" + "---|" * ncol)
                for row in rows:
                    out.append("| " + " | ".join(_md_cell(x) for x in row + [""] * (ncol - len(row))) + " |")
            out.append("")
        for p in (pictures(s) if pictures else []):
            out += [("[Bild]" if plain else "*[Bild: " + os.path.basename(p) + "]*"), ""]
        if s.get("notes"):
            out += [("Notiz: " if plain else "> ") + s["notes"], ""]
    return ("\n".join(out).strip() + "\n").encode("utf-8")


def txt(c, style=None, pictures=None, base=None):
    return md(c, style, pictures, base, plain=True)


def csv_(c, style=None, pictures=None, base=None):
    """Every table one after another (an empty line between), German Excel's ; for German documents."""
    b = io.StringIO()
    w = csv.writer(b, delimiter=";" if (c.get("lang") or "de") == "de" else ",", lineterminator="\r\n")
    first = True
    for s in _sections(c):
        tb = s.get("table")
        if not tb:
            continue
        if not first:
            w.writerow([])
        first = False
        for row in ([tb["head"]] if tb.get("head") else []) + (tb.get("rows") or []):
            w.writerow([csv_cell(x) for x in row])
    if first:   # no table: the headings and bullet points as one column
        for s in _sections(c):
            for x in [s.get("heading") or ""] + (s.get("bullets") or []) + (s.get("paras") or []):
                w.writerow([csv_cell(x)])
    return ("\ufeff" + b.getvalue()).encode("utf-8")


BUILD = {"docx": docx, "pptx": pptx, "xlsx": xlsx, "pdf": pdf, "md": md, "txt": txt, "csv": csv_}


def render(c, fmt, style_key="spark-hell", pictures=None, base=None):
    """The file in this format: bytes. pictures(section) -> [paths] (already checked), base: an own template."""
    if fmt not in BUILD:
        raise ValueError("format: " + ", ".join(FORMATS))
    st = style_of(style_key)
    if fmt not in ("docx", "pptx"):
        base = None
    return BUILD[fmt](c, st, pictures, base)


def german_date(y, m, d, lang="de"):
    if lang == "en":
        import calendar
        return f"{calendar.month_name[m]} {d}, {y}"
    return f"{d}. {MONTHS[m - 1]} {y}"
